"""Service for generating Word DOCX manuscripts."""
import os
import logging
from datetime import datetime
from sqlalchemy.orm import Session

from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.config import settings
from app.models.manuscript_state import ManuscriptState
from app.models.manuscript_section import ManuscriptSection
from app.models.citation import Citation
from app.models.export_record import ExportRecord
from app.models.export_template import ExportTemplate
from app.export.pre_export_validator import PreExportValidator

logger = logging.getLogger(__name__)


class DocxGeneratorService:
    """Generates Word DOCX files from manuscript data."""

    def generate(self, manuscript_id: int, template_id: int | None, db: Session) -> ExportRecord:
        """Generate a DOCX file for the manuscript."""
        # 1. Validate
        validator = PreExportValidator()
        validation = validator.validate(manuscript_id, db)

        manuscript = db.query(ManuscriptState).filter(
            ManuscriptState.id == manuscript_id
        ).first()
        if not manuscript:
            raise ValueError("Manuscript not found")

        # Create export record
        record = ExportRecord(
            project_id=manuscript.project_id,
            manuscript_state_id=manuscript_id,
            template_id=template_id,
            export_type="docx",
            status="generating",
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        try:
            if not validation.is_valid:
                raise RuntimeError(
                    "Validation failed: " + "; ".join(validation.errors)
                )

            # 2. Load template style
            style_config = self._default_style()
            if template_id:
                template = db.query(ExportTemplate).filter(
                    ExportTemplate.id == template_id
                ).first()
                if template and template.style_config:
                    style_config = template.style_config

            # 3. Load sections
            sections = (
                db.query(ManuscriptSection)
                .filter(ManuscriptSection.manuscript_state_id == manuscript_id)
                .order_by(ManuscriptSection.sort_order)
                .all()
            )

            # 4. Load citations
            citations = (
                db.query(Citation)
                .filter(Citation.project_id == manuscript.project_id)
                .all()
            )
            cite_key_to_num = {}
            for i, c in enumerate(citations, 1):
                cite_key_to_num[c.cite_key] = i

            # 5. Build document
            doc = Document()
            self._apply_page_margins(doc, style_config)
            self._render_title(doc, manuscript, style_config)
            self._render_abstract(doc, manuscript, style_config)

            # Track which outline paths have DB sections
            db_section_keys = {s.section_key for s in sections}

            for section in sections:
                self._render_section(doc, section, cite_key_to_num, style_config)

            # Render outline sections that have no DB counterpart (fallback)
            outline = manuscript.outline if isinstance(manuscript.outline, list) else []
            if outline:
                self._render_outline_fallback(doc, outline, db_section_keys, "", 1, style_config)

            self._render_references(doc, citations, style_config)
            self._render_ai_disclosure(doc, style_config)

            # 6. Save
            export_dir = os.path.join(settings.EXPORT_PATH, str(manuscript.project_id))
            os.makedirs(export_dir, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            file_name = f"{timestamp}_manuscript.docx"
            file_path = os.path.join(export_dir, file_name)
            doc.save(file_path)
            file_size = os.path.getsize(file_path)

            # 7. Update record
            record.file_path = file_path
            record.file_name = file_name
            record.file_size = file_size
            record.status = "completed"
            record.completed_at = datetime.now()

            # Add warnings if any
            if validation.warnings:
                record.error_message = "Warnings: " + "; ".join(validation.warnings)

        except Exception as exc:
            record.status = "failed"
            record.error_message = str(exc)
            logger.error(f"DOCX generation failed: {exc}")

        db.commit()
        db.refresh(record)
        return record

    def _default_style(self) -> dict:
        return {
            "page": {"margin_top": 2.54, "margin_bottom": 2.54, "margin_left": 3.18, "margin_right": 3.18},
            "title": {"font_name": "Times New Roman", "font_size": 16, "bold": True, "alignment": "center"},
            "heading1": {"font_name": "Times New Roman", "font_size": 14, "bold": True},
            "heading2": {"font_name": "Times New Roman", "font_size": 12, "bold": True},
            "body": {"font_name": "Times New Roman", "font_size": 12, "line_spacing": 1.5},
            "reference": {"font_name": "Times New Roman", "font_size": 10},
        }

    def _apply_page_margins(self, doc: Document, style: dict):
        page = style.get("page", {})
        for section in doc.sections:
            if "margin_top" in page:
                section.top_margin = Cm(page["margin_top"])
            if "margin_bottom" in page:
                section.bottom_margin = Cm(page["margin_bottom"])
            if "margin_left" in page:
                section.left_margin = Cm(page["margin_left"])
            if "margin_right" in page:
                section.right_margin = Cm(page["margin_right"])

    def _render_title(self, doc: Document, manuscript: ManuscriptState, style: dict):
        title_style = style.get("title", {})
        p = doc.add_paragraph()
        alignment = WD_ALIGN_PARAGRAPH.CENTER if title_style.get("alignment") == "center" else WD_ALIGN_PARAGRAPH.LEFT
        p.alignment = alignment
        run = p.add_run(manuscript.title or "Untitled")
        run.font.name = title_style.get("font_name", "Times New Roman")
        run.font.size = Pt(title_style.get("font_size", 16))
        run.bold = title_style.get("bold", True)

    def _render_abstract(self, doc: Document, manuscript: ManuscriptState, style: dict):
        if manuscript.abstract:
            h_style = style.get("heading1", {})
            heading = doc.add_heading("Abstract", level=1)
            for run in heading.runs:
                run.font.name = h_style.get("font_name", "Times New Roman")
                run.font.size = Pt(h_style.get("font_size", 14))

            b_style = style.get("body", {})
            p = doc.add_paragraph(manuscript.abstract)
            for run in p.runs:
                run.font.name = b_style.get("font_name", "Times New Roman")
                run.font.size = Pt(b_style.get("font_size", 12))

    def _render_section(self, doc: Document, section: ManuscriptSection,
                        cite_key_to_num: dict, style: dict):
        h_style = style.get("heading1", {})
        heading = doc.add_heading(section.title or section.section_key, level=1)
        for run in heading.runs:
            run.font.name = h_style.get("font_name", "Times New Roman")
            run.font.size = Pt(h_style.get("font_size", 14))

        content = section.generated_content or section.content or ""
        if content:
            # Replace citation keys with [N] format
            for key, num in cite_key_to_num.items():
                content = content.replace(f"\\cite{{{key}}}", f"[{num}]")
                content = content.replace(f"@{key}", f"[{num}]")

            b_style = style.get("body", {})
            for paragraph_text in content.split("\n\n"):
                paragraph_text = paragraph_text.strip()
                if paragraph_text:
                    p = doc.add_paragraph(paragraph_text)
                    for run in p.runs:
                        run.font.name = b_style.get("font_name", "Times New Roman")
                        run.font.size = Pt(b_style.get("font_size", 12))

        # Render tables from table_refs
        table_refs = section.table_refs if isinstance(section.table_refs, list) else []
        for tref in table_refs:
            caption = tref.get("caption", "")
            p = doc.add_paragraph(f"Table: {caption}")
            p.runs[0].bold = True if p.runs else False

        # Render figures
        figure_refs = section.figure_refs if isinstance(section.figure_refs, list) else []
        for fref in figure_refs:
            caption = fref.get("caption", "")
            file_path = fref.get("file_path", "")
            if file_path and os.path.isfile(file_path):
                try:
                    doc.add_picture(file_path, width=Inches(5.5))
                except Exception:
                    doc.add_paragraph(f"[Figure: {caption}]")
            else:
                doc.add_paragraph(f"[Figure: {caption}]")

    def _render_outline_fallback(self, doc: Document, nodes: list,
                                  db_section_keys: set, parent_path: str,
                                  depth: int, style: dict):
        """Recursively render outline nodes that have no DB section as fallback content."""
        for idx, node in enumerate(nodes):
            path = f"{parent_path}{idx}" if not parent_path else f"{parent_path}.{idx}"
            section_key = f"section_{path.replace('.', '_')}"

            # Determine heading level (depth 1 → heading 1, depth 2 → heading 2, etc.)
            heading_level = min(depth, 3)  # Cap at heading 3
            h_key = f"heading{heading_level}" if heading_level <= 2 else "heading2"
            h_style = style.get(h_key, style.get("heading1", {}))

            # Build display number (1-indexed)
            display_num = ".".join(str(int(n) + 1) for n in path.split("."))
            title = node.get("title") or node.get("section") or node.get("name") or f"第 {display_num} 节"

            if section_key not in db_section_keys:
                # No DB section — render outline as fallback content
                heading = doc.add_heading(title, level=heading_level)
                for run in heading.runs:
                    run.font.name = h_style.get("font_name", "Times New Roman")
                    run.font.size = Pt(h_style.get("font_size", 14 if heading_level == 1 else 12))

                b_style = style.get("body", {})

                # Render description
                desc = node.get("description", "")
                if desc:
                    p = doc.add_paragraph(desc)
                    for run in p.runs:
                        run.font.name = b_style.get("font_name", "Times New Roman")
                        run.font.size = Pt(b_style.get("font_size", 12))

                # Render key points
                key_points = node.get("key_points", [])
                if key_points and isinstance(key_points, list):
                    kp_heading = doc.add_paragraph()
                    run = kp_heading.add_run("关键要点：")
                    run.bold = True
                    run.font.name = b_style.get("font_name", "Times New Roman")
                    run.font.size = Pt(b_style.get("font_size", 12))
                    for kp in key_points:
                        p = doc.add_paragraph(kp, style="List Bullet")
                        for run in p.runs:
                            run.font.name = b_style.get("font_name", "Times New Roman")
                            run.font.size = Pt(b_style.get("font_size", 12))

                # Render writing plan — ONLY for leaf nodes (no children)
                children = node.get("sections") or node.get("subsections") or []
                has_children = bool(children and isinstance(children, list) and len(children) > 0)
                writing_plan = node.get("writing_plan", [])
                if not has_children and writing_plan and isinstance(writing_plan, list):
                    wp_heading = doc.add_paragraph()
                    run = wp_heading.add_run("写作计划：")
                    run.bold = True
                    run.font.name = b_style.get("font_name", "Times New Roman")
                    run.font.size = Pt(b_style.get("font_size", 12))
                    for i, wp in enumerate(writing_plan, 1):
                        topic = wp.get("paragraph_topic", "")
                        wc = wp.get("word_count", "")
                        refs = wp.get("key_references", [])
                        line = f"段落{i}: {topic}"
                        if wc:
                            line += f" (约{wc}字)"
                        if refs:
                            line += f"  参考: {', '.join(refs)}"
                        p = doc.add_paragraph(line)
                        for run in p.runs:
                            run.font.name = b_style.get("font_name", "Times New Roman")
                            run.font.size = Pt(b_style.get("font_size", 12))

            # Recurse into children regardless
            # Support both "sections" (nested objects) and "subsections" (flat strings)
            children = node.get("sections") or node.get("subsections") or []
            if children and isinstance(children[0], str):
                children = [{"title": s, "sections": [], "key_points": [], "description": ""} for s in children]
            if children and isinstance(children, list):
                self._render_outline_fallback(
                    doc, children, db_section_keys, path, depth + 1, style
                )

    def _render_references(self, doc: Document, citations: list, style: dict):
        if not citations:
            return
        h_style = style.get("heading1", {})
        heading = doc.add_heading("References", level=1)
        for run in heading.runs:
            run.font.name = h_style.get("font_name", "Times New Roman")
            run.font.size = Pt(h_style.get("font_size", 14))

        r_style = style.get("reference", {})
        for i, c in enumerate(citations, 1):
            text = f"[{i}] {c.authors}. {c.title}. "
            if c.venue:
                text += f"{c.venue}, "
            if c.year:
                text += f"{c.year}."
            if c.doi:
                text += f" DOI: {c.doi}"
            p = doc.add_paragraph(text)
            for run in p.runs:
                run.font.name = r_style.get("font_name", "Times New Roman")
                run.font.size = Pt(r_style.get("font_size", 10))

    def _render_ai_disclosure(self, doc: Document, style: dict):
        h_style = style.get("heading1", {})
        heading = doc.add_heading("AI Usage Disclosure", level=1)
        for run in heading.runs:
            run.font.name = h_style.get("font_name", "Times New Roman")
            run.font.size = Pt(h_style.get("font_size", 14))

        disclosure = (
            "This manuscript was prepared with AI-assisted tools for literature analysis, "
            "idea generation, and drafting support. All content has been reviewed and verified "
            "by the authors. The AI tools used are acknowledged as writing aids and are not "
            "listed as authors. All experimental results presented are based on actual data "
            "and have not been fabricated."
        )
        b_style = style.get("body", {})
        p = doc.add_paragraph(disclosure)
        for run in p.runs:
            run.font.name = b_style.get("font_name", "Times New Roman")
            run.font.size = Pt(b_style.get("font_size", 12))
