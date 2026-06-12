"""Service for generating response letters to reviewer comments."""
import os
import logging
from datetime import datetime
from sqlalchemy.orm import Session

from docx import Document
from docx.shared import Pt

from app.config import settings
from app.models.manuscript_state import ManuscriptState
from app.models.reviewer_simulation import ReviewerSimulation
from app.models.export_record import ExportRecord
from app.services.llm_service import get_llm_service

logger = logging.getLogger(__name__)


class ResponseLetterService:
    """Generates point-by-point response letters to reviewer simulations."""

    def generate(
        self,
        manuscript_id: int,
        additional_context: str,
        db: Session,
    ) -> ExportRecord:
        manuscript = db.query(ManuscriptState).filter(
            ManuscriptState.id == manuscript_id
        ).first()
        if not manuscript:
            raise ValueError("Manuscript not found")

        record = ExportRecord(
            project_id=manuscript.project_id,
            manuscript_state_id=manuscript_id,
            export_type="response_letter",
            status="generating",
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        try:
            # Load reviewer simulations
            simulations = (
                db.query(ReviewerSimulation)
                .filter(ReviewerSimulation.project_id == manuscript.project_id)
                .all()
            )

            if not simulations:
                raise ValueError(
                    "No reviewer simulations found. Generate reviewer simulations first."
                )

            # Get LLM to generate responses
            llm = get_llm_service(settings)
            
            if hasattr(llm, 'generate_response_letter'):
                response_content = llm.generate_response_letter(
                    manuscript_title=manuscript.title or "Untitled",
                    simulations=simulations,
                    additional_context=additional_context,
                )
            else:
                response_content = self._fallback_response(manuscript, simulations)

            # Build DOCX
            doc = Document()
            
            title_style = doc.add_heading("Response to Reviewers", level=1)
            doc.add_paragraph(f"Manuscript: {manuscript.title or 'Untitled'}")
            doc.add_paragraph("")

            body_text = response_content if isinstance(response_content, str) else response_content.get("content", "")
            for paragraph in body_text.split("\n\n"):
                paragraph = paragraph.strip()
                if paragraph:
                    # Check if it's a reviewer heading
                    if paragraph.startswith("## "):
                        doc.add_heading(paragraph[3:], level=2)
                    elif paragraph.startswith("### "):
                        doc.add_heading(paragraph[4:], level=3)
                    elif paragraph.startswith("**") and paragraph.endswith("**"):
                        p = doc.add_paragraph(paragraph[2:-2])
                        for run in p.runs:
                            run.bold = True
                    else:
                        p = doc.add_paragraph(paragraph)
                        for run in p.runs:
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(12)

            # Save
            export_dir = os.path.join(settings.EXPORT_PATH, str(manuscript.project_id))
            os.makedirs(export_dir, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            file_name = f"{timestamp}_response_letter.docx"
            file_path = os.path.join(export_dir, file_name)
            doc.save(file_path)
            file_size = os.path.getsize(file_path)

            record.file_path = file_path
            record.file_name = file_name
            record.file_size = file_size
            record.status = "completed"
            record.completed_at = datetime.now()

        except Exception as exc:
            record.status = "failed"
            record.error_message = str(exc)
            logger.error(f"Response letter generation failed: {exc}")

        db.commit()
        db.refresh(record)
        return record

    def _fallback_response(self, manuscript, simulations) -> str:
        parts = [f"We thank the reviewers for their constructive feedback on our manuscript \"{manuscript.title or 'Untitled'}\".\n\n"]
        
        for sim in simulations:
            role = sim.reviewer_role or "Reviewer"
            parts.append(f"## {role} ({sim.expertise_area or ''})\n\n")
            parts.append(f"**Overall Assessment:** {sim.overall_assessment or 'N/A'}\n\n")
            
            # Major issues
            major = sim.major_issues if isinstance(sim.major_issues, list) else []
            if major:
                parts.append("### Major Issues\n\n")
                for i, issue in enumerate(major, 1):
                    parts.append(f"**Issue {i}:** {issue}\n\n")
                    parts.append(f"**Response:** [Author response to be filled]\n\n")
            
            # Minor issues
            minor = sim.minor_issues if isinstance(sim.minor_issues, list) else []
            if minor:
                parts.append("### Minor Issues\n\n")
                for i, issue in enumerate(minor, 1):
                    parts.append(f"**Issue {i}:** {issue}\n\n")
                    parts.append(f"**Response:** [Author response to be filled]\n\n")

            # Suggested experiments
            suggested = sim.suggested_experiments if isinstance(sim.suggested_experiments, list) else []
            if suggested:
                parts.append("### Suggested Experiments\n\n")
                for i, exp in enumerate(suggested, 1):
                    parts.append(f"**Suggestion {i}:** {exp}\n\n")
                    parts.append(f"**Response:** [Author response to be filled]\n\n")

        return "\n".join(parts)
