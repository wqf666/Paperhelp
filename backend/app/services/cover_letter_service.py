"""Service for generating cover letters."""
import os
import logging
from datetime import datetime
from sqlalchemy.orm import Session

from docx import Document
from docx.shared import Pt

from app.config import settings
from app.models.manuscript_state import ManuscriptState
from app.models.export_record import ExportRecord
from app.services.llm_service import get_llm_service

logger = logging.getLogger(__name__)


class CoverLetterService:
    """Generates cover letter DOCX for journal/conference submissions."""

    def generate(
        self,
        manuscript_id: int,
        recipient_name: str,
        recipient_title: str,
        journal_or_conference: str,
        additional_notes: str,
        db: Session,
    ) -> ExportRecord:
        manuscript = db.query(ManuscriptState).filter(
            ManuscriptState.id == manuscript_id
        ).first()
        if not manuscript:
            raise ValueError("Manuscript not found")

        record = ExportRecord(
            project_id=manuscript.project_id,
            manuscript_state_id=manuscript_id,
            export_type="cover_letter",
            status="generating",
        )
        db.add(record)
        db.commit()
        db.refresh(record)

        try:
            # Get LLM to generate cover letter content
            llm = get_llm_service(settings)
            
            if hasattr(llm, 'generate_cover_letter'):
                letter_content = llm.generate_cover_letter(
                    manuscript_title=manuscript.title or "Untitled",
                    manuscript_abstract=manuscript.abstract or "",
                    recipient_name=recipient_name,
                    recipient_title=recipient_title,
                    journal_or_conference=journal_or_conference,
                    additional_notes=additional_notes,
                )
            else:
                # Fallback template
                letter_content = self._fallback_template(
                    manuscript, recipient_name, recipient_title,
                    journal_or_conference, additional_notes
                )

            # Build DOCX
            doc = Document()
            
            # Date
            doc.add_paragraph(datetime.now().strftime("%B %d, %Y"))
            doc.add_paragraph("")
            
            # Recipient
            doc.add_paragraph(f"Dear {recipient_title} {recipient_name},")
            doc.add_paragraph("")

            # Body
            body_text = letter_content if isinstance(letter_content, str) else letter_content.get("content", "")
            for paragraph in body_text.split("\n\n"):
                paragraph = paragraph.strip()
                if paragraph:
                    p = doc.add_paragraph(paragraph)
                    for run in p.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(12)

            # Closing
            doc.add_paragraph("")
            doc.add_paragraph("Sincerely,")
            doc.add_paragraph("[Author Names]")

            # Save
            export_dir = os.path.join(settings.EXPORT_PATH, str(manuscript.project_id))
            os.makedirs(export_dir, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            file_name = f"{timestamp}_cover_letter.docx"
            file_path = os.path.join(export_dir, file_name)
            doc.save(file_path)
            file_size = os.path.getsize(file_path)

            record.file_path = file_path
            record.file_name = file_name
            record.file_size = file_size
            record.status = "completed"
            record.completed_at = datetime.now()

        except Exception as exc:
            record.status = "failed"
            record.error_message = str(exc)
            logger.error(f"Cover letter generation failed: {exc}")

        db.commit()
        db.refresh(record)
        return record

    def _fallback_template(self, manuscript, recipient_name, recipient_title, journal, notes) -> str:
        title = manuscript.title or "Our Manuscript"
        return (
            f"I am writing to submit our manuscript entitled \"{title}\" for consideration "
            f"in {journal or 'your journal/conference'}.\n\n"
            f"{manuscript.abstract or 'This manuscript presents our research findings.'}\n\n"
            f"We believe this work is well-suited for {journal or 'your venue'} because it "
            f"addresses important challenges in the field and presents novel contributions.\n\n"
            f"{('Additional notes: ' + notes) if notes else ''}\n\n"
            f"We confirm that this manuscript has not been published elsewhere and is not "
            f"under consideration by another journal or conference."
        )
