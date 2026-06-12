"""Service for converting Word documents to LaTeX source."""
import re
import logging

logger = logging.getLogger(__name__)

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class WordToLatexConverter:
    """Convert Word DOCX files to basic LaTeX source."""

    def convert(self, file_content: bytes, file_name: str) -> dict:
        """Convert a Word document to LaTeX.
        
        Returns:
            dict with keys:
            - latex_source: str (the .tex content)
            - warnings: list[str] (conversion warnings)
        """
        if not HAS_DOCX:
            return {"latex_source": "% python-docx not available", "warnings": ["python-docx not installed"]}

        import io
        warnings = []
        
        try:
            doc = Document(io.BytesIO(file_content))
        except Exception as exc:
            return {"latex_source": f"% Failed to parse document: {exc}", "warnings": [str(exc)]}

        body_parts = []
        in_list = False
        list_type = None  # "itemize" or "enumerate"

        for para in doc.paragraphs:
            style_name = para.style.name if para.style else "Normal"
            text = self._convert_inline(para)

            if not text.strip() and style_name == "Normal":
                if in_list:
                    body_parts.append(f"\\end{{{list_type}}}")
                    in_list = False
                    list_type = None
                body_parts.append("")
                continue

            # Headings
            if style_name.startswith("Heading"):
                if in_list:
                    body_parts.append(f"\\end{{{list_type}}}")
                    in_list = False
                    list_type = None
                level = self._heading_level(style_name)
                cmd = ["\\section", "\\subsection", "\\subsubsection", "\\paragraph"][min(level - 1, 3)]
                body_parts.append(f"{cmd}{{{text}}}")
                body_parts.append("")
                continue

            # List items
            if style_name.startswith("List"):
                if "Bullet" in style_name or "Unordered" in style_name:
                    target_type = "itemize"
                else:
                    target_type = "enumerate"
                
                if not in_list or list_type != target_type:
                    if in_list:
                        body_parts.append(f"\\end{{{list_type}}}")
                    body_parts.append(f"\\begin{{{target_type}}}")
                    in_list = True
                    list_type = target_type
                body_parts.append(f"  \\item {text}")
                continue

            # Normal paragraph
            if in_list:
                body_parts.append(f"\\end{{{list_type}}}")
                in_list = False
                list_type = None
            body_parts.append(text)
            body_parts.append("")

        # Close any remaining list
        if in_list:
            body_parts.append(f"\\end{{{list_type}}}")

        # Process tables
        for i, table in enumerate(doc.tables):
            table_latex = self._convert_table(table, i + 1)
            body_parts.append("")
            body_parts.append(table_latex)

        # Assemble full document
        body_text = "\n".join(body_parts)
        
        latex_source = (
            "\\documentclass[11pt]{article}\n"
            "\\usepackage[utf8]{inputenc}\n"
            "\\usepackage{graphicx}\n"
            "\\usepackage{amsmath}\n"
            "\\usepackage{hyperref}\n\n"
            "\\title{Converted from " + self._escape_latex(file_name) + "}\n"
            "\\date{}\n\n"
            "\\begin{document}\n"
            "\\maketitle\n\n"
            f"{body_text}\n\n"
            "\\end{document}\n"
        )

        if warnings:
            logger.info(f"Word to LaTeX conversion warnings: {warnings}")

        return {"latex_source": latex_source, "warnings": warnings}

    def _heading_level(self, style_name: str) -> int:
        """Extract heading level from style name like 'Heading 1'."""
        match = re.search(r"(\d+)", style_name)
        return int(match.group(1)) if match else 1

    def _convert_inline(self, paragraph) -> str:
        """Convert inline formatting of a paragraph to LaTeX."""
        parts = []
        for run in paragraph.runs:
            text = self._escape_latex(run.text)
            if not text:
                continue
            if run.bold and run.italic:
                text = f"\\textbf{{\\textit{{{text}}}}}"
            elif run.bold:
                text = f"\\textbf{{{text}}}"
            elif run.italic:
                text = f"\\textit{{{text}}}"
            parts.append(text)
        return "".join(parts)

    def _escape_latex(self, text: str) -> str:
        """Escape special LaTeX characters."""
        if not text:
            return ""
        special = {
            "&": "\\&", "%": "\\%", "$": "\\$", "#": "\\#",
            "_": "\\_", "{": "\\{", "}": "\\}",
            "~": "\\textasciitilde{}", "^": "\\textasciicircum{}",
        }
        result = text
        # Replace backslash first (before adding new ones)
        result = result.replace("\\", "\\textbackslash{}")
        for char, replacement in special.items():
            result = result.replace(char, replacement)
        return result

    def _convert_table(self, table, table_num: int) -> str:
        """Convert a Word table to LaTeX tabular."""
        rows = []
        for row in table.rows:
            cells = [self._escape_latex(cell.text.strip()) for cell in row.cells]
            rows.append(cells)

        if not rows:
            return f"% Empty table {table_num}"

        col_count = max(len(r) for r in rows) if rows else 1
        col_spec = "|" + "|".join(["c"] * col_count) + "|"

        lines = [
            f"\\begin{{table}}[htbp]",
            f"\\centering",
            f"\\caption{{Table {table_num}}}",
            f"\\label{{tab:{table_num}}}",
            f"\\begin{{tabular}}{{{col_spec}}}",
            "\\hline",
        ]
        for row in rows:
            # Pad row to match column count
            while len(row) < col_count:
                row.append("")
            lines.append(" & ".join(row) + " \\\\")
            lines.append("\\hline")
        lines.extend([
            "\\end{tabular}",
            "\\end{table}",
        ])
        return "\n".join(lines)
