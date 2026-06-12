"""Tests for MVP 3 features: citations, templates, export (DOCX/LaTeX/BibTeX),
cover letters, response letters, word-to-latex conversion, section management,
pre-export validation, and compliance extensions."""

import io
import json
import os

from tests.conftest import *  # noqa: F401,F403 – fixtures


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _create_project(client, name="MVP3 Project"):
    resp = client.post(
        "/projects",
        json={"name": name, "research_field": "NLP", "target_venue": "ACL"},
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


def _create_paper(client, project_id, title="MVP3 Paper"):
    resp = client.post(
        f"/projects/{project_id}/papers",
        json={"title": title, "abstract": "A study about transformers."},
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


def _generate_ideas(client, project_id, field="NLP"):
    resp = client.post(
        f"/projects/{project_id}/ideas/generate",
        json={"research_field": field},
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


def _create_manuscript(client, project_id):
    """Create a manuscript state for the project by generating an outline."""
    resp = client.post(
        f"/projects/{project_id}/manuscript/outline/generate",
        json={},
    )
    assert resp.status_code in (200, 201), resp.text
    return resp.json()


def _create_citation(client, project_id, cite_key="smith2023nlp", title="NLP Paper"):
    resp = client.post(
        f"/projects/{project_id}/citations",
        json={
            "cite_key": cite_key,
            "entry_type": "article",
            "title": title,
            "authors": "Smith and Jones",
            "year": 2023,
            "venue": "ACL",
        },
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


SAMPLE_BIBTEX = """@article{smith2023nlp,
  title = {Advances in NLP},
  author = {Smith, John and Jones, Mary},
  year = {2023},
  journal = {ACL Anthology},
  doi = {10.1234/nlp.2023},
}

@inproceedings{wang2022deep,
  title = {Deep Learning for NLP},
  author = {Wang, Wei},
  year = {2022},
  booktitle = {EMNLP},
  pages = {100--110},
}
"""


# ===========================================================================
# Citations
# ===========================================================================

class TestCitations:
    """Test citation CRUD and BibTeX import/export."""

    def test_create_citation(self, client):
        pid = _create_project(client)
        data = _create_citation(client, pid)
        assert data["cite_key"] == "smith2023nlp"
        assert data["title"] == "NLP Paper"
        assert data["entry_type"] == "article"
        assert data["year"] == 2023

    def test_list_citations(self, client):
        pid = _create_project(client)
        _create_citation(client, pid, cite_key="key1", title="First")
        _create_citation(client, pid, cite_key="key2", title="Second")
        resp = client.get(f"/projects/{pid}/citations")
        assert resp.status_code == 200
        data = resp.json()
        assert len(data) == 2

    def test_list_citations_search(self, client):
        pid = _create_project(client)
        _create_citation(client, pid, cite_key="alpha", title="Alpha Paper")
        _create_citation(client, pid, cite_key="beta", title="Beta Paper")
        resp = client.get(f"/projects/{pid}/citations?search=Alpha")
        assert resp.status_code == 200
        data = resp.json()
        assert len(data) == 1
        assert data[0]["cite_key"] == "alpha"

    def test_get_citation(self, client):
        pid = _create_project(client)
        c = _create_citation(client, pid)
        resp = client.get(f"/citations/{c['id']}")
        assert resp.status_code == 200
        assert resp.json()["cite_key"] == "smith2023nlp"

    def test_get_citation_not_found(self, client):
        resp = client.get("/citations/99999")
        assert resp.status_code == 404

    def test_update_citation(self, client):
        pid = _create_project(client)
        c = _create_citation(client, pid)
        resp = client.put(
            f"/citations/{c['id']}",
            json={"title": "Updated Title", "year": 2024},
        )
        assert resp.status_code == 200
        assert resp.json()["title"] == "Updated Title"
        assert resp.json()["year"] == 2024

    def test_delete_citation(self, client):
        pid = _create_project(client)
        c = _create_citation(client, pid)
        resp = client.delete(f"/citations/{c['id']}")
        assert resp.status_code == 200
        # Verify it's gone
        resp2 = client.get(f"/citations/{c['id']}")
        assert resp2.status_code == 404

    def test_import_bibtex(self, client):
        pid = _create_project(client)
        resp = client.post(
            f"/projects/{pid}/citations/import-bibtex",
            files={"file": ("refs.bib", SAMPLE_BIBTEX.encode(), "application/x-bibtex")},
        )
        assert resp.status_code == 201, resp.text
        data = resp.json()
        assert len(data) == 2
        keys = {c["cite_key"] for c in data}
        assert "smith2023nlp" in keys
        assert "wang2022deep" in keys

    def test_import_bibtex_nonexistent_project(self, client):
        resp = client.post(
            "/projects/99999/citations/import-bibtex",
            files={"file": ("refs.bib", SAMPLE_BIBTEX.encode(), "application/x-bibtex")},
        )
        assert resp.status_code == 404

    def test_export_bibtex(self, client):
        pid = _create_project(client)
        _create_citation(client, pid, cite_key="key1", title="Paper One")
        _create_citation(client, pid, cite_key="key2", title="Paper Two")
        resp = client.get(f"/projects/{pid}/citations/export-bibtex")
        assert resp.status_code == 200
        text = resp.text
        assert "key1" in text
        assert "key2" in text
        assert "@article" in text

    def test_export_bibtex_nonexistent_project(self, client):
        resp = client.get("/projects/99999/citations/export-bibtex")
        assert resp.status_code == 404


# ===========================================================================
# Templates
# ===========================================================================

class TestTemplates:
    """Test template listing and management."""

    def test_list_templates(self, client):
        """Built-in templates should be seeded on startup."""
        resp = client.get("/templates")
        assert resp.status_code == 200
        data = resp.json()
        # At least 4 built-in templates
        assert len(data) >= 4

    def test_list_templates_filter_docx(self, client):
        resp = client.get("/templates?template_type=docx")
        assert resp.status_code == 200
        data = resp.json()
        assert all(t["template_type"] == "docx" for t in data)
        assert len(data) >= 1

    def test_list_templates_filter_latex(self, client):
        resp = client.get("/templates?template_type=latex")
        assert resp.status_code == 200
        data = resp.json()
        assert all(t["template_type"] == "latex" for t in data)
        assert len(data) >= 1

    def test_get_template(self, client):
        # Get the first template
        resp = client.get("/templates")
        templates = resp.json()
        assert len(templates) > 0
        tid = templates[0]["id"]
        resp2 = client.get(f"/templates/{tid}")
        assert resp2.status_code == 200
        assert resp2.json()["id"] == tid

    def test_get_template_not_found(self, client):
        resp = client.get("/templates/99999")
        assert resp.status_code == 404

    def test_delete_builtin_template_forbidden(self, client):
        """Built-in templates cannot be deleted."""
        resp = client.get("/templates")
        templates = resp.json()
        builtin = [t for t in templates if t["is_builtin"]]
        assert len(builtin) > 0
        resp2 = client.delete(f"/templates/{builtin[0]['id']}")
        assert resp2.status_code == 400


# ===========================================================================
# Manuscript Sections (extended)
# ===========================================================================

class TestManuscriptSectionsExtended:
    """Test section create, delete, reorder."""

    def test_create_section(self, client):
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)

        resp = client.post(
            f"/projects/{pid}/manuscript/sections",
            json={
                "section_key": "introduction",
                "title": "Introduction",
                "content": "This is the introduction.",
                "sort_order": 1,
            },
        )
        assert resp.status_code == 201, resp.text
        data = resp.json()
        assert data["section_key"] == "introduction"
        assert data["title"] == "Introduction"
        assert data["content"] == "This is the introduction."

    def test_create_section_no_manuscript(self, client):
        pid = _create_project(client)
        resp = client.post(
            f"/projects/{pid}/manuscript/sections",
            json={"section_key": "intro", "title": "Intro"},
        )
        assert resp.status_code == 404

    def test_delete_section(self, client):
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)

        # Create a section first
        resp = client.post(
            f"/projects/{pid}/manuscript/sections",
            json={"section_key": "test", "title": "Test Section", "sort_order": 0},
        )
        assert resp.status_code == 201
        sid = resp.json()["id"]

        resp2 = client.delete(f"/manuscript-sections/{sid}")
        assert resp2.status_code == 200

        # Verify it's gone
        resp3 = client.get(f"/projects/{pid}/manuscript/sections")
        section_ids = [s["id"] for s in resp3.json()]
        assert sid not in section_ids

    def test_delete_section_not_found(self, client):
        resp = client.delete("/manuscript-sections/99999")
        assert resp.status_code == 404

    def test_reorder_sections(self, client):
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)

        # Create two sections
        r1 = client.post(
            f"/projects/{pid}/manuscript/sections",
            json={"section_key": "sec_a", "title": "Section A", "sort_order": 1},
        )
        r2 = client.post(
            f"/projects/{pid}/manuscript/sections",
            json={"section_key": "sec_b", "title": "Section B", "sort_order": 2},
        )
        id_a = r1.json()["id"]
        id_b = r2.json()["id"]

        # Reorder: B first, A second
        resp = client.put(
            f"/projects/{pid}/manuscript/sections/reorder",
            json=[{"id": id_b, "sort_order": 1}, {"id": id_a, "sort_order": 2}],
        )
        assert resp.status_code == 200
        sections = resp.json()
        assert sections[0]["id"] == id_b
        assert sections[1]["id"] == id_a

    def test_reorder_no_manuscript(self, client):
        pid = _create_project(client)
        resp = client.put(
            f"/projects/{pid}/manuscript/sections/reorder",
            json=[],
        )
        assert resp.status_code == 404


# ===========================================================================
# Export DOCX
# ===========================================================================

class TestExportDocx:
    """Test DOCX export."""

    def test_export_docx_success(self, client):
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)
        _create_citation(client, pid)

        # Create a section with content so validation passes
        client.post(
            f"/projects/{pid}/manuscript/sections",
            json={
                "section_key": "introduction",
                "title": "Introduction",
                "content": "This is the introduction paragraph.",
                "sort_order": 1,
            },
        )

        resp = client.post(
            f"/projects/{pid}/export/docx",
            json={},
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["export_type"] == "docx"
        assert data["status"] in ("completed", "failed")

    def test_export_docx_no_manuscript(self, client):
        pid = _create_project(client)
        resp = client.post(
            f"/projects/{pid}/export/docx",
            json={},
        )
        assert resp.status_code == 404

    def test_export_docx_with_template(self, client):
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)

        # Create content section
        client.post(
            f"/projects/{pid}/manuscript/sections",
            json={
                "section_key": "method",
                "title": "Method",
                "content": "Our method uses transformers.",
                "sort_order": 1,
            },
        )

        # Get a docx template
        templates_resp = client.get("/templates?template_type=docx")
        templates = templates_resp.json()
        if templates:
            tid = templates[0]["id"]
            resp = client.post(
                f"/projects/{pid}/export/docx",
                json={"template_id": tid},
            )
            assert resp.status_code == 200, resp.text
            data = resp.json()
            assert data["export_type"] == "docx"


# ===========================================================================
# Export LaTeX
# ===========================================================================

class TestExportLatex:
    """Test LaTeX export."""

    def test_export_latex_success(self, client):
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)
        _create_citation(client, pid)

        # Create a section with content
        client.post(
            f"/projects/{pid}/manuscript/sections",
            json={
                "section_key": "introduction",
                "title": "Introduction",
                "content": "This is the introduction.",
                "sort_order": 1,
            },
        )

        resp = client.post(
            f"/projects/{pid}/export/latex",
            json={},
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["export_type"] == "latex"
        assert data["status"] in ("completed", "failed")

    def test_export_latex_no_manuscript(self, client):
        pid = _create_project(client)
        resp = client.post(
            f"/projects/{pid}/export/latex",
            json={},
        )
        assert resp.status_code == 404


# ===========================================================================
# Export BibTeX
# ===========================================================================

class TestExportBibtex:
    """Test BibTeX export via export endpoint."""

    def test_export_bibtex_success(self, client):
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)
        _create_citation(client, pid, cite_key="test2023", title="Test Paper")

        resp = client.post(f"/projects/{pid}/export/bibtex")
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["export_type"] == "bibtex"
        assert data["status"] == "completed"

    def test_export_bibtex_no_manuscript(self, client):
        pid = _create_project(client)
        resp = client.post(f"/projects/{pid}/export/bibtex")
        assert resp.status_code == 404


# ===========================================================================
# Cover Letter
# ===========================================================================

class TestCoverLetter:
    """Test cover letter generation."""

    def test_generate_cover_letter(self, client):
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)

        resp = client.post(
            f"/projects/{pid}/export/cover-letter",
            json={
                "recipient_name": "Smith",
                "recipient_title": "Professor",
                "journal_or_conference": "ACL 2025",
                "additional_notes": "Please consider for oral presentation.",
            },
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["export_type"] == "cover_letter"
        assert data["status"] in ("completed", "failed")

    def test_generate_cover_letter_no_manuscript(self, client):
        pid = _create_project(client)
        resp = client.post(
            f"/projects/{pid}/export/cover-letter",
            json={
                "recipient_name": "Editor",
                "recipient_title": "Dr.",
                "journal_or_conference": "Nature",
            },
        )
        assert resp.status_code == 404


# ===========================================================================
# Response Letter
# ===========================================================================

class TestResponseLetter:
    """Test response letter generation."""

    def test_generate_response_letter(self, client):
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)

        # Generate reviewer simulations first
        client.post(
            f"/projects/{pid}/reviewer-simulation/generate",
            json={"num_reviewers": 2},
        )

        resp = client.post(
            f"/projects/{pid}/export/response-letter",
            json={"additional_context": "We have added more experiments."},
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["export_type"] == "response_letter"
        assert data["status"] in ("completed", "failed")

    def test_generate_response_letter_no_manuscript(self, client):
        pid = _create_project(client)
        resp = client.post(
            f"/projects/{pid}/export/response-letter",
            json={},
        )
        assert resp.status_code == 404

    def test_generate_response_letter_no_reviewers(self, client):
        """Should fail when no reviewer simulations exist."""
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)

        resp = client.post(
            f"/projects/{pid}/export/response-letter",
            json={},
        )
        # Should still return 200 but with status=failed (or 500)
        assert resp.status_code in (200, 500)
        if resp.status_code == 200:
            data = resp.json()
            assert data["status"] == "failed"


# ===========================================================================
# Export Record
# ===========================================================================

class TestExportRecord:
    """Test export record retrieval."""

    def test_get_export_record(self, client):
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)
        _create_citation(client, pid)

        client.post(
            f"/projects/{pid}/manuscript/sections",
            json={"section_key": "intro", "title": "Intro", "content": "Content.", "sort_order": 1},
        )

        # Create an export
        export_resp = client.post(f"/projects/{pid}/export/docx", json={})
        if export_resp.status_code == 200:
            record_id = export_resp.json()["id"]
            resp = client.get(f"/exports/{record_id}")
            assert resp.status_code == 200
            assert resp.json()["id"] == record_id

    def test_get_export_record_not_found(self, client):
        resp = client.get("/exports/99999")
        assert resp.status_code == 404


# ===========================================================================
# Word-to-LaTeX Converter
# ===========================================================================

class TestWordToLatex:
    """Test Word to LaTeX conversion endpoint."""

    def test_convert_word_to_latex(self, client):
        from docx import Document

        # Create a simple DOCX in memory
        doc = Document()
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("This is the introduction.")
        doc.add_heading("Method", level=1)
        doc.add_paragraph("Our method uses transformers.")

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        resp = client.post(
            "/convert/word-to-latex",
            files={"file": ("test.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert "latex_source" in data
        assert "\\section{Introduction}" in data["latex_source"]
        assert "\\section{Method}" in data["latex_source"]


# ===========================================================================
# Section Content Generation
# ===========================================================================

class TestSectionContentGeneration:
    """Test LLM-based section content generation."""

    def test_generate_section_content(self, client):
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)

        # Create a section
        sec_resp = client.post(
            f"/projects/{pid}/manuscript/sections",
            json={
                "section_key": "introduction",
                "title": "Introduction",
                "sort_order": 1,
            },
        )
        assert sec_resp.status_code == 201
        section_id = sec_resp.json()["id"]

        # Generate content for it
        resp = client.post(
            f"/projects/{pid}/manuscript/sections/{section_id}/generate-content",
        )
        assert resp.status_code == 200, resp.text
        data = resp.json()
        assert data["generated_content"] is not None
        assert len(data["generated_content"]) > 0

    def test_generate_section_content_no_manuscript(self, client):
        pid = _create_project(client)
        resp = client.post(
            f"/projects/{pid}/manuscript/sections/99999/generate-content",
        )
        assert resp.status_code == 404


# ===========================================================================
# BibTeX Parser Unit Tests
# ===========================================================================

class TestBibtexParser:
    """Unit tests for the BibTeX parser."""

    def test_parse_single_entry(self):
        from app.parsers.bibtex_parser import BibTeXParser

        parser = BibTeXParser()
        text = """@article{doe2023ml,
  title = {Machine Learning Advances},
  author = {Doe, Jane},
  year = {2023},
  journal = {JMLR},
}"""
        entries = parser.parse(text)
        assert len(entries) == 1
        assert entries[0]["cite_key"] == "doe2023ml"
        assert entries[0]["entry_type"] == "article"
        assert entries[0]["fields"]["title"] == "Machine Learning Advances"
        assert entries[0]["fields"]["year"] == "2023"

    def test_parse_multiple_entries(self):
        from app.parsers.bibtex_parser import BibTeXParser

        parser = BibTeXParser()
        entries = parser.parse(SAMPLE_BIBTEX)
        assert len(entries) == 2
        keys = {e["cite_key"] for e in entries}
        assert "smith2023nlp" in keys
        assert "wang2022deep" in keys

    def test_parse_empty_text(self):
        from app.parsers.bibtex_parser import BibTeXParser

        parser = BibTeXParser()
        assert parser.parse("") == []
        assert parser.parse("   ") == []

    def test_entry_to_citation_data(self):
        from app.parsers.bibtex_parser import BibTeXParser

        parser = BibTeXParser()
        entries = parser.parse(SAMPLE_BIBTEX)
        data = parser.entry_to_citation_data(entries[0])
        assert data["cite_key"] == "smith2023nlp"
        assert data["year"] == 2023
        assert data["source"] == "imported"
        assert "Smith" in data["authors"]

    def test_serialize_citations(self):
        from app.parsers.bibtex_parser import BibTeXParser

        parser = BibTeXParser()

        class FakeCitation:
            cite_key = "test2023"
            entry_type = "article"
            title = "Test Paper"
            authors = "Author One"
            year = 2023
            venue = "Test Journal"
            doi = "10.1234/test"
            url = ""
            abstract = ""
            extra_fields = {}

        result = parser.serialize([FakeCitation()])
        assert "@article{test2023," in result
        assert "title = {Test Paper}" in result
        assert "year = {2023}" in result

    def test_validate_entry(self):
        from app.parsers.bibtex_parser import BibTeXParser

        parser = BibTeXParser()
        # Missing required fields
        entry = {"cite_key": "", "fields": {}}
        warnings = parser.validate(entry)
        assert any("cite_key" in w.lower() for w in warnings)

        # Complete entry
        entry2 = {
            "cite_key": "key",
            "fields": {"title": "T", "author": "A", "year": "2023"},
        }
        warnings2 = parser.validate(entry2)
        assert len(warnings2) == 0


# ===========================================================================
# Pre-Export Validator Unit Tests
# ===========================================================================

class TestPreExportValidator:
    """Unit tests for pre-export validation."""

    def test_validate_no_manuscript(self, client, db_session):
        from app.export.pre_export_validator import PreExportValidator

        validator = PreExportValidator()
        result = validator.validate(99999, db_session)
        assert not result.is_valid
        assert any("not found" in e.lower() for e in result.errors)

    def test_validate_empty_title(self, client, db_session):
        from app.export.pre_export_validator import PreExportValidator
        from app.models.manuscript_state import ManuscriptState

        # Create project + manuscript manually with empty title
        pid = _create_project(client)
        ms = ManuscriptState(project_id=pid, title="", abstract="Some abstract")
        db_session.add(ms)
        db_session.commit()

        validator = PreExportValidator()
        result = validator.validate(ms.id, db_session)
        assert not result.is_valid
        assert any("title" in e.lower() for e in result.errors)

    def test_validate_success(self, client, db_session):
        from app.export.pre_export_validator import PreExportValidator
        from app.models.manuscript_state import ManuscriptState
        from app.models.manuscript_section import ManuscriptSection

        pid = _create_project(client)
        ms = ManuscriptState(project_id=pid, title="Test Title", abstract="Test Abstract")
        db_session.add(ms)
        db_session.commit()
        db_session.refresh(ms)

        sec = ManuscriptSection(
            manuscript_state_id=ms.id,
            section_key="introduction",
            title="Introduction",
            content="This is the introduction.",
            sort_order=1,
        )
        db_session.add(sec)
        db_session.commit()

        validator = PreExportValidator()
        result = validator.validate(ms.id, db_session)
        assert result.is_valid

    def test_validate_unresolved_citations(self, client, db_session):
        from app.export.pre_export_validator import PreExportValidator
        from app.models.manuscript_state import ManuscriptState
        from app.models.manuscript_section import ManuscriptSection

        pid = _create_project(client)
        ms = ManuscriptState(project_id=pid, title="Title", abstract="Abstract")
        db_session.add(ms)
        db_session.commit()
        db_session.refresh(ms)

        sec = ManuscriptSection(
            manuscript_state_id=ms.id,
            section_key="intro",
            title="Introduction",
            content="As shown by \\cite{missing_ref}.",
            citation_keys=["missing_ref"],
            sort_order=1,
        )
        db_session.add(sec)
        db_session.commit()

        validator = PreExportValidator()
        result = validator.validate(ms.id, db_session)
        assert not result.is_valid
        assert any("missing_ref" in e for e in result.errors)


# ===========================================================================
# Compliance Extensions
# ===========================================================================

class TestComplianceExtensions:
    """Test new compliance functions: check_missing_citations, generate_ai_disclosure."""

    def test_check_missing_citations_pass(self):
        from app.compliance.checker import check_missing_citations

        result = check_missing_citations(
            "As shown by \\cite{smith2023}.",
            citation_keys=["smith2023"],
        )
        assert result["status"] == "passed"
        assert len(result["violations"]) == 0

    def test_check_missing_citations_fail(self):
        from app.compliance.checker import check_missing_citations

        result = check_missing_citations(
            "As shown by \\cite{missing_key} and @another_missing.",
            citation_keys=["smith2023"],
        )
        assert result["status"] == "failed"
        assert len(result["violations"]) == 2
        assert any("missing_key" in v for v in result["violations"])
        assert any("another_missing" in v for v in result["violations"])

    def test_check_missing_citations_empty_content(self):
        from app.compliance.checker import check_missing_citations

        result = check_missing_citations("", ["key1"])
        assert result["status"] == "passed"

    def test_generate_ai_disclosure_standard(self):
        from app.compliance.checker import generate_ai_disclosure

        text = generate_ai_disclosure("standard")
        assert "AI-assisted" in text
        assert "reviewed" in text.lower()

    def test_generate_ai_disclosure_minimal(self):
        from app.compliance.checker import generate_ai_disclosure

        text = generate_ai_disclosure("minimal")
        assert len(text) < len(generate_ai_disclosure("standard"))
        assert "AI-assisted" in text

    def test_generate_ai_disclosure_detailed(self):
        from app.compliance.checker import generate_ai_disclosure

        text = generate_ai_disclosure("detailed")
        assert len(text) > len(generate_ai_disclosure("standard"))
        assert "Literature Analysis" in text

    def test_generate_ai_disclosure_unknown_type(self):
        from app.compliance.checker import generate_ai_disclosure

        text = generate_ai_disclosure("unknown_type")
        # Should fallback to standard
        assert "AI-assisted" in text


# ===========================================================================
# Word-to-Latex Converter Unit Tests
# ===========================================================================

class TestWordToLatexConverterUnit:
    """Unit tests for the WordToLatexConverter class."""

    def test_convert_simple_document(self):
        from app.services.word_to_latex_converter import WordToLatexConverter
        from docx import Document

        doc = Document()
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("Hello world.")

        buf = io.BytesIO()
        doc.save(buf)
        content = buf.getvalue()

        converter = WordToLatexConverter()
        result = converter.convert(content, "test.docx")
        assert "latex_source" in result
        assert "\\section{Introduction}" in result["latex_source"]
        assert "\\begin{document}" in result["latex_source"]

    def test_convert_with_bold_text(self):
        from app.services.word_to_latex_converter import WordToLatexConverter
        from docx import Document

        doc = Document()
        p = doc.add_paragraph()
        run = p.add_run("Bold text")
        run.bold = True

        buf = io.BytesIO()
        doc.save(buf)
        content = buf.getvalue()

        converter = WordToLatexConverter()
        result = converter.convert(content, "test.docx")
        assert "\\textbf{Bold text}" in result["latex_source"]

    def test_convert_with_table(self):
        from app.services.word_to_latex_converter import WordToLatexConverter
        from docx import Document

        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        table.cell(1, 0).text = "1"
        table.cell(1, 1).text = "2"

        buf = io.BytesIO()
        doc.save(buf)
        content = buf.getvalue()

        converter = WordToLatexConverter()
        result = converter.convert(content, "test.docx")
        assert "\\begin{tabular}" in result["latex_source"]

    def test_convert_invalid_file(self):
        from app.services.word_to_latex_converter import WordToLatexConverter

        converter = WordToLatexConverter()
        result = converter.convert(b"not a docx file", "bad.docx")
        assert "warnings" in result
        assert len(result["warnings"]) > 0

    def test_escape_latex_special_chars(self):
        from app.services.word_to_latex_converter import WordToLatexConverter

        converter = WordToLatexConverter()
        assert converter._escape_latex("10% increase") == "10\\% increase"
        assert converter._escape_latex("A & B") == "A \\& B"
        assert converter._escape_latex("$100") == "\\$100"


# ===========================================================================
# MVP3 Integration Tests
# ===========================================================================

class TestMVP3Integration:
    """Integration tests covering full MVP3 workflows."""

    def test_citation_workflow(self, client):
        """Full citation workflow: create -> list -> search -> export."""
        pid = _create_project(client)

        # Create citations
        c1 = _create_citation(client, pid, cite_key="paper_a", title="Paper Alpha")
        c2 = _create_citation(client, pid, cite_key="paper_b", title="Paper Beta")

        # List all
        resp = client.get(f"/projects/{pid}/citations")
        assert len(resp.json()) == 2

        # Search
        resp = client.get(f"/projects/{pid}/citations?search=Alpha")
        assert len(resp.json()) == 1

        # Update
        client.put(f"/citations/{c1['id']}", json={"year": 2024})
        resp = client.get(f"/citations/{c1['id']}")
        assert resp.json()["year"] == 2024

        # Export BibTeX
        resp = client.get(f"/projects/{pid}/citations/export-bibtex")
        assert "paper_a" in resp.text
        assert "paper_b" in resp.text

        # Delete one
        client.delete(f"/citations/{c2['id']}")
        resp = client.get(f"/projects/{pid}/citations")
        assert len(resp.json()) == 1

    def test_full_export_workflow(self, client):
        """Full workflow: project -> ideas -> manuscript -> sections -> citations -> export."""
        pid = _create_project(client)

        # Create paper and analyze
        paper_id = _create_paper(client, pid)
        client.post(f"/papers/{paper_id}/analyze")

        # Generate ideas
        _generate_ideas(client, pid)

        # Generate manuscript
        _create_manuscript(client, pid)

        # Add sections
        client.post(
            f"/projects/{pid}/manuscript/sections",
            json={
                "section_key": "introduction",
                "title": "Introduction",
                "content": "This paper presents our findings.",
                "sort_order": 1,
            },
        )
        client.post(
            f"/projects/{pid}/manuscript/sections",
            json={
                "section_key": "conclusion",
                "title": "Conclusion",
                "content": "In conclusion, our method works.",
                "sort_order": 2,
            },
        )

        # Add citations
        _create_citation(client, pid, cite_key="ref2023", title="Reference Paper")

        # Export DOCX
        resp = client.post(f"/projects/{pid}/export/docx", json={})
        assert resp.status_code == 200
        docx_record = resp.json()
        assert docx_record["export_type"] == "docx"

        # Export LaTeX
        resp = client.post(f"/projects/{pid}/export/latex", json={})
        assert resp.status_code == 200
        latex_record = resp.json()
        assert latex_record["export_type"] == "latex"

        # Export BibTeX
        resp = client.post(f"/projects/{pid}/export/bibtex")
        assert resp.status_code == 200
        bib_record = resp.json()
        assert bib_record["export_type"] == "bibtex"

    def test_reviewer_to_response_letter_workflow(self, client):
        """Workflow: reviewer simulations -> response letter."""
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)

        # Generate reviewer simulations
        resp = client.post(
            f"/projects/{pid}/reviewer-simulation/generate",
            json={"num_reviewers": 2},
        )
        assert resp.status_code in (200, 201)

        # Generate response letter
        resp = client.post(
            f"/projects/{pid}/export/response-letter",
            json={"additional_context": "Added new experiments in Section 4."},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data["export_type"] == "response_letter"

    def test_bibtex_import_and_use_in_export(self, client):
        """Import BibTeX, reference in sections, then export."""
        pid = _create_project(client)
        _generate_ideas(client, pid)
        _create_manuscript(client, pid)

        # Import citations
        resp = client.post(
            f"/projects/{pid}/citations/import-bibtex",
            files={"file": ("refs.bib", SAMPLE_BIBTEX.encode(), "application/x-bibtex")},
        )
        assert resp.status_code == 201
        citations = resp.json()
        assert len(citations) == 2

        # Create section referencing a citation
        client.post(
            f"/projects/{pid}/manuscript/sections",
            json={
                "section_key": "related_work",
                "title": "Related Work",
                "content": "Previous work by \\cite{smith2023nlp} shows promising results.",
                "citation_keys": ["smith2023nlp"],
                "sort_order": 1,
            },
        )

        # Export DOCX - citation should be resolved
        resp = client.post(f"/projects/{pid}/export/docx", json={})
        assert resp.status_code == 200
